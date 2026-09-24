"""Create a portable source ZIP, with a redacted copy of the reference workbook."""

import argparse
import hashlib
import json
import os
import sys
from copy import copy
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path, PurePosixPath
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile

ROOT = Path(__file__).resolve().parents[1]
TREES = (
    "backend/app",
    "backend/migrations",
    "backend/tests",
    "frontend/src",
    "frontend/mailbox",
    "frontend/public",
    "frontend/tests",
    "frontend/docs",
    "data",
    "docs",
    "plan",
    "scripts",
    "usecases",
)
FILES = (
    "README.md",
    ".gitignore",
    ".env.example",
    "backend/pyproject.toml",
    "backend/requirements.lock",
    "backend/openapi.json",
    "frontend/package.json",
    "frontend/package-lock.json",
    "frontend/index.html",
    "frontend/main.tsx",
    "frontend/tsconfig.json",
    "frontend/playwright.config.ts",
    "frontend/vite.config.ts",
    "frontend/vite.mailbox.config.ts",
    "aaa.png",
    "outlook.png",
    "outlook-inbox.png",
)
EXCLUDED_DIRS = {
    ".git",
    ".agents",
    ".codex",
    ".claude",
    ".venv",
    ".local",
    "node_modules",
    "__pycache__",
    ".pytest_cache",
    ".ruff_cache",
    "htmlcov",
    "dist",
    "dist-mailbox",
    "test-results",
    "playwright-report",
    "releases",
}
EXCLUDED_SUFFIXES = {
    ".pyc",
    ".pyo",
    ".log",
    ".db",
    ".sqlite",
    ".sqlite3",
    ".sqlite3-shm",
    ".sqlite3-wal",
    ".tsbuildinfo",
    ".onetoc2",
    ".zip",
    ".pem",
    ".key",
    ".pfx",
    ".p12",
    ".bak",
    ".tmp",
}
WORKBOOK = "Inquiry_Responses_Extracted_Pro.xlsx"


def excluded(path):
    parts = path.parts
    return (
        any(p in EXCLUDED_DIRS or p.endswith(".egg-info") for p in parts)
        or path.suffix.lower() in EXCLUDED_SUFFIXES
        or path.name in {".DS_Store", "Thumbs.db", ".coverage"}
        or path.name.startswith("~$")
        or (path.name.startswith(".env") and path.name != ".env.example")
    )


def require_local_file(path):
    if not path.resolve().is_relative_to(ROOT) or not path.is_file():
        raise ValueError(f"Missing or external package input: {path.relative_to(ROOT)}")
    for part in (path, *path.parents):
        if part == ROOT:
            break
        if part.is_symlink() or part.is_junction():
            raise ValueError(
                "Links and junctions cannot be included in a source package."
            )


def source_files():
    entries = {name: ROOT / name for name in FILES}
    for tree in TREES:
        base = ROOT / tree
        if not base.is_dir():
            raise ValueError(f"Required source directory is missing: {tree}")
        for directory, folders, filenames in os.walk(base, followlinks=False):
            folders[:] = [
                name
                for name in folders
                if not excluded(Path(directory, name).relative_to(ROOT))
            ]
            for name in filenames:
                path = Path(directory, name)
                relative = path.relative_to(ROOT)
                if not excluded(relative):
                    entries[relative.as_posix()] = path
    for path in entries.values():
        require_local_file(path)
    return entries


def json_bytes(value):
    return (json.dumps(value, indent=2, ensure_ascii=False) + "\n").encode("utf-8")


def redacted_sources():
    """Export values only, removing known credentials and internal sender settings.

    Reconcile against the exported workbook so check-data also works after extraction.
    The original workbook and versioned artifacts are never changed.
    """
    try:
        from openpyxl import Workbook, load_workbook

        sys.path.insert(0, str(ROOT / "backend"))
        from app.source_catalog import SECRET_MARKER, reconcile_sources
    except ImportError as exc:
        raise RuntimeError(
            "Run scripts/setup.ps1 first, then package using the project's Python environment."
        ) from exc

    original = load_workbook(ROOT / "docs" / WORKBOOK, data_only=True, keep_links=False)
    shared = Workbook()
    shared.remove(shared.active)
    redactions = []
    redacted_labels = {}
    try:
        for sheet in original:
            output = shared.create_sheet(sheet.title)
            output.sheet_state = sheet.sheet_state
            output.freeze_panes = sheet.freeze_panes
            for key, dimension in sheet.column_dimensions.items():
                output.column_dimensions[key].width = dimension.width
            for row in sheet:
                number = row[0].row
                sensitive = (
                    sheet.title == "Extracted Data" and 474 <= number <= 483
                ) or any(
                    isinstance(cell.value, str) and SECRET_MARKER.search(cell.value)
                    for cell in row
                )
                if sensitive and sheet.title == "Extracted Data":
                    # Preserve duplicate-group structure without exporting original labels.
                    label_key = (sheet.title, row[0].value, row[1].value)
                    group = redacted_labels.setdefault(
                        label_key, len(redacted_labels) + 1
                    )
                for cell in row:
                    value = cell.value
                    if sensitive:
                        value = None
                        if sheet.title == "Extracted Data":
                            if cell.column == 1:
                                value = (
                                    cell.value
                                )  # Preserve the source category and physical row.
                            elif cell.column == 2:
                                value = f"Redacted source group {group}"
                            elif cell.column == 3:
                                value = "Removed from this shared copy: credentials or internal sender configuration."
                        elif cell.column == 1:
                            value = "Redacted from the shared source copy."
                    destination = output.cell(cell.row, cell.column, value)
                    if cell.has_style:
                        destination.font = copy(cell.font)
                        destination.fill = copy(cell.fill)
                        destination.border = copy(cell.border)
                        destination.alignment = copy(cell.alignment)
                        destination.number_format = cell.number_format
                if sensitive:
                    redactions.append({"sheet": sheet.title, "row": number})
            for cells in sheet.merged_cells.ranges:
                output.merge_cells(str(cells))
        shared.properties.creator = "Correspondence source export"
        buffer = BytesIO()
        shared.save(buffer)
        content = buffer.getvalue()
    finally:
        original.close()
        shared.close()

    staging = ROOT / ".local" / "source-package-staging"
    if not staging.resolve().is_relative_to(ROOT):
        raise ValueError(
            "The temporary export directory must stay inside the workspace."
        )
    staging.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix="export-", dir=staging) as directory:
        temporary = Path(directory)
        (temporary / WORKBOOK).write_bytes(content)
        (temporary / "Class - Sub Class.xlsx").write_bytes(
            (ROOT / "docs/Class - Sub Class.xlsx").read_bytes()
        )
        reconciled = reconcile_sources(temporary)
    overrides = {
        f"docs/{WORKBOOK}": content,
        "data/knowledge/taxonomy.v1.json": json_bytes(reconciled["taxonomy"]),
        "data/knowledge/source_manifest.v1.json": json_bytes(reconciled["manifest"]),
        "data/knowledge/reconciliation.v1.json": json_bytes(reconciled["report"]),
    }
    return overrides, redactions


def verify_archive(path):
    with ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("ZIP integrity check failed.")
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise ValueError("ZIP contains duplicate paths.")
        manifest = json.loads(archive.read("PACKAGE_MANIFEST.json"))
        if set(names) != set(manifest) | {"PACKAGE_MANIFEST.json"}:
            raise ValueError("ZIP contents differ from the manifest.")
        for name, digest in manifest.items():
            relative = PurePosixPath(name)
            if (
                relative.is_absolute()
                or ".." in relative.parts
                or ":" in name
                or "\\" in name
                or excluded(Path(name))
            ):
                raise ValueError(f"Disallowed archive path: {name}")
            if hashlib.sha256(archive.read(name)).hexdigest() != digest:
                raise ValueError(f"ZIP file hash mismatch: {name}")
    return len(manifest)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        help="Destination .zip file; existing files are never overwritten.",
    )
    parser.add_argument(
        "--verify", type=Path, help="Verify an existing package without rebuilding it."
    )
    args = parser.parse_args()
    if args.verify:
        print(
            json.dumps(
                {
                    "verified": str(args.verify.resolve()),
                    "files": verify_archive(args.verify),
                }
            )
        )
        return
    version = json.loads((ROOT / "frontend/package.json").read_text(encoding="utf-8"))[
        "version"
    ]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    destination = (
        args.output
        or ROOT / ".local/releases" / f"correspondence-source-{version}-{stamp}.zip"
    ).resolve()
    if destination.suffix.lower() != ".zip":
        raise ValueError("The output filename must end in .zip.")
    if destination.exists():
        raise FileExistsError(
            "The destination already exists. Choose a new ZIP filename."
        )
    entries = source_files()
    overrides, redactions = redacted_sources()
    overrides["START_HERE.md"] = (ROOT / "plan/Source_Sharing_Guide.md").read_bytes()
    overrides["PACKAGE_INFO.json"] = json_bytes(
        {
            "package_type": "source",
            "version": version,
            "built_at_utc": stamp,
            "original_source_documents_included": True,
            "reference_workbook": "Values-only copy with credential and sender rows redacted",
            "redacted_rows": redactions,
            "local_env_included": False,
            "runtime_case_data_included": False,
            "installed_dependencies_included": False,
        }
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    manifest = {}
    created = False
    try:
        with ZipFile(
            destination, "x", compression=ZIP_DEFLATED, compresslevel=9
        ) as archive:
            created = True
            for name in sorted(set(entries) | set(overrides)):
                content = (
                    overrides[name] if name in overrides else entries[name].read_bytes()
                )
                manifest[name] = hashlib.sha256(content).hexdigest()
                archive.writestr(name, content)
            archive.writestr("PACKAGE_MANIFEST.json", json_bytes(manifest))
        file_count = verify_archive(destination)
    except Exception:
        if created:
            destination.unlink(
                missing_ok=True
            )  # Remove only the incomplete ZIP created here.
        raise
    report = {
        "path": str(destination),
        "files": file_count,
        "size_bytes": destination.stat().st_size,
        "sha256": hashlib.sha256(destination.read_bytes()).hexdigest(),
        "verified": True,
        "redacted_reference_rows": len(redactions),
    }
    receipt = destination.with_suffix(".zip.sha256")
    receipt.write_text(f"{report['sha256']}  {destination.name}\n", encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
