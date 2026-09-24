"""Offline, explicit source reconciliation. Raw response text is never exported."""

import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from app.config import WORKSPACE_ROOT
from app.fixture_contracts import KnowledgeFixture

DATA_ROOT = WORKSPACE_ROOT / "data"
RESPONSE_FILE = "Inquiry_Responses_Extracted_Pro.xlsx"
CATEGORY_RANGES = [
    (2, 23, "1098-1099"),
    (24, 107, "Simple Responses"),
    (108, 155, "Account-Ownership Updates"),
    (156, 163, "Attorney Correspondence"),
    (164, 210, "Escrow"),
    (211, 289, "PMT - STMT"),
    (290, 330, "Taxes - Insurance"),
    (331, 336, "BANA-Not QWR's - Links"),
    (337, 368, "Credit -FC-LM-FEMA"),
    (369, 386, "Assumption"),
    (387, 408, "Forms-REO-Job Aids-FAQ"),
    (409, 458, "Website"),
    (459, 473, "Recast"),
    (474, 483, "Signitures-Kitework Emails"),
    (484, 515, "Task"),
    (516, 528, "HELOC-SOCURE"),
]
SECRET_MARKER = re.compile(
    r"(?i)(password|passwd|pwd|api[_ -]?key|secret|access[_ -]?token)\s*[:=]\s*\S+"
)
REVIEW_REASONS = {
    19: "heading_body_mismatch",
    54: "unrelated_topics_merged",
    229: "mixed_instructions_and_frequency_conflict",
    292: "merged_incomplete_tax_wording",
    515: "task_and_unrelated_identity_instructions",
    519: "incomplete_heloc_fragment",
    29: "production_bankruptcy_disclosure_requires_policy_review",
    34: "production_bankruptcy_disclosure_requires_policy_review",
    523: "heloc_only_procedure_outside_fixed_mortgage_scope",
}


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def classify_source_row(number: int, fields: tuple, curated_keys: list[str]) -> tuple[str, str]:
    # Known credential-bearing row is rejected before examining or serializing text.
    if number == 474:
        return "excluded", "credential_bearing_source_row"
    if 475 <= number <= 483:
        return "excluded", "production_sender_contact_configuration"
    joined = " ".join(str(value or "") for value in fields)
    if SECRET_MARKER.search(joined):
        return "excluded", "possible_credential_assignment"
    if number in REVIEW_REASONS:
        return "requires_review", REVIEW_REASONS[number]
    body = str(fields[2] or "").strip()
    if len(body) < 50:
        return "requires_review", "short_or_fragmentary_entry"
    if curated_keys:
        return "curated", "selected_guidance_rewritten_for_synthetic_demo"
    return "requires_review", "not_curated_for_selected_release"


def reconcile_sources(source_dir: Path = WORKSPACE_ROOT / "docs") -> dict:
    curated = [
        KnowledgeFixture.model_validate(row)
        for row in read_json(DATA_ROOT / "knowledge/curated.v1.json")
    ]
    taxonomy_path = source_dir / "Class - Sub Class.xlsx"
    workbook = load_workbook(taxonomy_path, read_only=True, data_only=True)
    try:
        rows = list(workbook["Sheet1"].iter_rows(values_only=True))
        if tuple(rows[0][:3]) != ("Work Type", "Class", "Sub Class"):
            raise ValueError("Taxonomy source headers changed; review the source snapshot.")
        taxonomy = [
            {
                "id": f"S03:r{number}",
                "work_type": row[0],
                "class_name": row[1],
                "subclass": row[2],
                "source_reference": f"S03:r{number}",
            }
            for number, row in enumerate(rows[1:], 2)
            if any(row)
        ]
    finally:
        workbook.close()
    combinations = {(row["work_type"], row["class_name"], row["subclass"]) for row in taxonomy}
    if (
        len(taxonomy) != 149
        or len(combinations) != 149
        or any(not all(key) for key in combinations)
    ):
        raise ValueError("Taxonomy must reconcile to 149 complete parent-qualified combinations.")

    response_path = source_dir / RESPONSE_FILE
    workbook = load_workbook(response_path, read_only=True, data_only=True)
    manifest = []
    label_groups = {}
    body_groups = {}
    try:
        sheet = workbook["Extracted Data"]
        iterator = sheet.iter_rows(values_only=True)
        if tuple(next(iterator)[:3]) != ("Category", "Sub Category", "Description"):
            raise ValueError("Knowledge source headers changed; review the source snapshot.")
        for number, fields in enumerate(iterator, 2):
            category = next(
                (label for start, end, label in CATEGORY_RANGES if start <= number <= end), None
            )
            if category is None or fields[0] != category or not all(fields[:3]):
                raise ValueError(
                    f"Source row {number} does not match the reconciled category/row layout."
                )
            reference = f"S04:r{number}"
            keys = [item.key for item in curated if reference in item.source_references]
            status, reason = classify_source_row(number, fields, keys)
            if status == "excluded" and keys:
                raise ValueError(f"Excluded source row {number} cannot supply curated content.")
            # No source subcategory, body, snippets, credential hashes or quarantine copies.
            manifest.append(
                {
                    "id": reference,
                    "row_number": number,
                    "category": category,
                    "status": status,
                    "reason": reason,
                    "curated_item_keys": keys,
                    "review_flags": [] if status == "curated" else [reason],
                }
            )
            if number != 474:
                label_groups.setdefault((category, str(fields[1])), []).append(number)
            if status != "excluded":
                normalized = " ".join(str(fields[2]).split())
                body_groups.setdefault(normalized, []).append(number)
        summary_has_525 = any(
            "525" in str(value)
            for row in workbook["Summary"].iter_rows(values_only=True)
            for value in row
            if value
        )
    finally:
        workbook.close()
    if len(manifest) != 527 or not summary_has_525:
        raise ValueError("Response source counts differ from the approved 527/525 reconciliation.")
    for groups, flag in (
        (label_groups, "repeated_category_subcategory"),
        (body_groups, "repeated_body_or_fragment"),
    ):
        for numbers in groups.values():
            if len(numbers) > 1:
                for number in numbers:
                    manifest[number - 2]["review_flags"].append(flag)
    report = {
        "version": 1,
        "taxonomy_count": len(taxonomy),
        "knowledge_row_count": len(manifest),
        "taxonomy_distinct_subclass_labels": len({row["subclass"] for row in taxonomy}),
        "summary_reported_rows": 525,
        "count_basis": "Extracted Data physical rows 2 through 528",
        "categories": dict(Counter(row["category"] for row in manifest)),
        "statuses": dict(Counter(row["status"] for row in manifest)),
        "curated_item_count": len(curated),
        "repeated_label_groups": sum(len(group) > 1 for group in label_groups.values()),
        "repeated_body_groups": sum(len(group) > 1 for group in body_groups.values()),
        "source_files": {
            taxonomy_path.name: sha256(taxonomy_path),
            response_path.name: sha256(response_path),
        },
        "raw_response_text_exported": False,
    }
    return {"taxonomy": taxonomy, "manifest": manifest, "report": report}


def export_catalogs() -> dict:
    result = reconcile_sources()
    for key, filename in (
        ("taxonomy", "taxonomy.v1.json"),
        ("manifest", "source_manifest.v1.json"),
        ("report", "reconciliation.v1.json"),
    ):
        write_json(DATA_ROOT / "knowledge" / filename, result[key])
    return result["report"]
